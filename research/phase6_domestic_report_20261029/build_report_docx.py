"""Create the editable report and a researcher-drawn logic figure."""
from pathlib import Path
import re,json
from PIL import Image, ImageDraw, ImageFont
import math
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent

def logic_figure():
    scale=250;w,h=3000,2000
    im=Image.new('RGB',(w,h),'white');dr=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/simsun.ttc',56)
    titlefont=ImageFont.truetype('C:/Windows/Fonts/simsun.ttc',64)
    smallfont=ImageFont.truetype('C:/Windows/Fonts/simsun.ttc',52)
    def xy(x,y):return (int(x*scale),int((8-y)*scale))
    def center(x,y,text,f=font):
        ls=text.split('\n');step=67;total=len(ls)*step
        for j,line in enumerate(ls):
            bb=dr.textbbox((0,0),line,font=f)
            dr.text((x-(bb[2]-bb[0])/2,y-total/2+j*step),line,font=f,fill='black')
    def box(x,y,bw,bh,text):
        x0,y0=xy(x,y+bh);x1,y1=xy(x+bw,y)
        dr.rounded_rectangle((x0,y0,x1,y1),radius=14,fill='#f3f6f8',outline='#66747e',width=3)
        center((x0+x1)/2,(y0+y1)/2,text)
    def arrow(a,b,label=None):
        a=xy(*a);b=xy(*b);dr.line([a,b],fill='#4c5963',width=4)
        angle=math.atan2(b[1]-a[1],b[0]-a[0]);length=22
        pts=[b,(b[0]-length*math.cos(angle-.45),b[1]-length*math.sin(angle-.45)),(b[0]-length*math.cos(angle+.45),b[1]-length*math.sin(angle+.45))]
        dr.polygon(pts,fill='#4c5963')
        if label:center((a[0]+b[0])/2+28,(a[1]+b[1])/2,label,smallfont)
    box(4.1,6.9,3.8,.65,'重要购买理由失去支持')
    box(3.4,5.4,5.2,.9,'审查重要性、公开联系\n可调整性与相称负担')
    box(.25,5.3,2.55,1.3,'规范前提\n程序性自主\n协助角色责任')
    box(9.2,5.2,2.55,1.3,'条件不足或不明\n处理事实与不确定性\n不猜测目标动机')
    box(3.85,3.75,4.3,.95,'普通更正是否已完成\n关联说明与可用修改机会')
    box(.6,1.95,4.75,1,'已完成\n承接选择，不增加形式确认')
    box(6.6,1.95,4.75,1,'尚有缺口\n简短补充，保留修改和结束')
    box(2.6,.25,6.8,1.0,'本人有效表达和授权决定预算\n拒绝与退出限制再次干预')
    arrow((6,6.9),(6,6.35));arrow((2.8,5.85),(3.35,5.85));arrow((8.65,5.85),(9.15,5.85))
    arrow((6,5.35),(6,4.75));arrow((4.7,3.7),(3,3.02),'是');arrow((7.3,3.7),(9,3.02),'否')
    arrow((3,1.9),(4.9,1.3));arrow((9,1.9),(7.1,1.3))
    center(1500,65,'从技术变化到规范判断及可实施动作',titlefont)
    im.save(ROOT/'figures/图1_论证框架.png',dpi=(300,300))

def setfont(run,name='SimSun',size=11.5,bold=False):
    run.font.name=name;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),name)

def runs(p,text,size=11.5,bold=False):
    text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',text)
    text=re.sub(r'另行测量。([\d\[\]]+)',lambda m:'另行\u2060测\u2060量\u2060。\u2060'+m.group(1),text)
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',text):
        if not part:continue
        isbold=part.startswith('**');t=part[2:-2] if isbold else part.strip('`')
        setfont(p.add_run(t),size=size,bold=bold or isbold)

def add_table(doc,lines):
    data=[[c.strip() for c in line.strip().strip('|').split('|')] for line in lines if not re.match(r'^\|[\s:|\-]+\|?$',line.strip())]
    n=len(data[0]);table=doc.add_table(rows=1,cols=n);table.autofit=False
    weights={3:[1.25,1.9,3.65],4:[1.15,1.65,2.5,1.5],5:[2.15,1.15,1.4,1.4,1.2],6:[1.75,.6,1.05,1.2,1.1,1.4]}[n]
    widths=[6.85*w/sum(weights) for w in weights]
    for i,w in enumerate(widths):table.columns[i].width=Inches(w)
    for rownum,rowdata in enumerate(data):
        row=table.rows[0] if rownum==0 else table.add_row()
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if rownum==0:trpr.append(OxmlElement('w:tblHeader'))
        for i,(cell,text) in enumerate(zip(row.cells,rowdata)):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();border=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');border.append(el)
            tcpr.append(border)
            mar=OxmlElement('w:tcMar')
            for edge,val in [('top','70'),('bottom','70'),('left','85'),('right','85')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF0F4' if rownum==0 else ('F8F9FA' if rownum%2==0 else 'FFFFFF'));tcpr.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.line_spacing=1.12;p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (rownum==0 or n==5 and i>0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next=rownum==0;runs(p,text,size=9.5,bold=rownum==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def build():
    logic_figure();doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.82);sec.right_margin=Inches(.82)
    sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
    for n in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[n];st.font.name='SimSun';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'SimSun')
        for bd in st._element.findall('.//'+qn('w:pBdr')):bd.getparent().remove(bd)
    normal=doc.styles['Normal'];normal.font.size=Pt(11.5);normal.paragraph_format.line_spacing=1.35;normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.first_line_indent=Inches(.29)
    for n,size in [('Title',20),('Heading 1',14),('Heading 2',12)]:
        st=doc.styles[n];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.first_line_indent=Inches(0);st.paragraph_format.space_before=Pt(12 if n!='Title' else 0);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
    h=sec.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;runs(h,'AI导购理由撤回后的预算重审',size=9)
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.paragraph_format.first_line_indent=Inches(0)
    runs(f,'第 ',size=9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field);runs(f,' 页',size=9)
    text=(ROOT/'draft.md').read_text(encoding='utf-8');lines=text.splitlines();i=0;refs=False;in_appendix=False
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
            add_table(doc,block);continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;runs(p,line[2:],size=20,bold=True)
        elif line.startswith('## '):
            name=line[3:].replace('、',' ')
            refs=name=='参考文献'
            in_appendix=name=='附录'
            heading=doc.add_heading(name,level=1)
            if name in ['参考文献','附录']:heading.paragraph_format.page_break_before=True
        elif line.startswith('### '):
            name=re.sub(r'^\d+\.\d+\s+','',line[4:]).replace('、',' ');doc.add_heading(name,level=2)
        elif line.startswith('!['):
            target=re.search(r'\]\(([^)]+)\)',line).group(1)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(ROOT/target),width=Inches(6.65))
        elif line.startswith('$$'):
            raise ValueError('Equation block needs an explicit renderer in this report.')
        else:
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            if line=='研究报告国内数据修订稿':p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);runs(p,line,size=10)
            elif re.match(r'^[表图]\d+',line):
                p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=line.startswith('表');runs(p,line,size=9.5,bold=line.startswith('表'))
            elif refs:
                p.paragraph_format.first_line_indent=Inches(-.25);p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.line_spacing=1.02;p.paragraph_format.space_after=Pt(1);runs(p,line,size=9)
            elif line.startswith('关键词：') or line.startswith('注：'):
                p.paragraph_format.first_line_indent=Inches(0);runs(p,line,size=10)
            else:
                if in_appendix:p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                runs(p,line)
        i+=1
    doc.core_properties.title='AI导购理由撤回后的预算重审';doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.subject='AI伦理研究报告国内数据修订稿';doc.core_properties.comments=''
    target=ROOT/'AI导购理由撤回后的预算重审_国内数据修订稿.docx';doc.save(target)
    print(json.dumps({'docx':target.name,'bytes':target.stat().st_size,'tables':len(doc.tables),'images':len(doc.inline_shapes)},ensure_ascii=False))

if __name__=='__main__':build()
