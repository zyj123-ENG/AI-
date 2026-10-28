from pathlib import Path
import csv, json, hashlib
from decimal import Decimal, ROUND_HALF_UP
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parent
sim=root.parent/'simulation_goal_repair_20261028'
def read(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
S=read(sim/'results/all_summary.csv'); C=read(sim/'results/paired_link_contrasts.csv')
D=read(sim/'results/base_diagnostics.csv')
ix={(r['scenario'],r['rule'],r['policy'],int(r['type'])):r for r in S}
ci={(r['scenario'],r['rule'],int(r['control']),int(r['type'])):r for r in C}
di={(r['rule'],r['policy'],int(r['type'])):r for r in D}
def get(p='D1_R1_C1',sc='base'):return ix[sc,'bounded_default',p,1]
def fmt(v,places=2,percent=False):
    x=Decimal(str(v))*(100 if percent else 1)
    return format(x.quantize(Decimal(1).scaleb(-places),rounding=ROUND_HALF_UP),'f')
def rate(r,k):return fmt(r[k],percent=True)
b=get('D1_R0_C1');n=get();d=ci['base','bounded_default',1,1]
drop=-100*float(d['delta_dependent_unreviewed_upgrade']);se=100*float(d['mcse_dependent_unreviewed_upgrade'])
increase=100*float(d['delta_goal_reviewed']);bad=100*float(d['delta_noncompatible_purchase'])
core=['D1_R0_C0','D1_R1_C0','D1_R0_C1','D1_R1_C1']
labels=['仅事实更正','增加关联说明','增加重选操作','关联说明与重选操作']
table1=[['处理条件','完成重评 %','未重评升级 %','请求后越限 %','不适配购买 %','处理负担']]
for p,label in zip(core,labels):
    r=get(p);table1.append([label,rate(r,'goal_reviewed'),rate(r,'dependent_unreviewed_upgrade'),rate(r,'budget_violation_after_request'),rate(r,'noncompatible_purchase'),fmt(r['processing_load'],3)])
table2=[['边界条件','模型中的比较结果','解释范围'],
 ['完全知情参考','各方案的选择与任务效用结果相同','来自参考规则设定'],
 ['更正未被理解','未重评升级均为100%','该模型中理解是重评前提'],
 ['关联线索效力为零','基准成本下主指标配对差为0','额外线索效力是待校准假设'],
 ['零商业偏向','效用差32.111增至34.837','重评增加不保证任务结果改善'],
 ['信息过载','负担4.916增至7.628','主指标仅变化0.06个百分点'],
 ['强选择性再确认','与所提方案结果完全相同','同构等参下的结构检查']]

blocks=[]
def title(t):blocks.append(('title',t))
def h(t,l=1):blocks.append(('h'+str(l),t))
def p(t):blocks.append(('p',t))
def eq(which):blocks.append(('eq',which))
def tab(cap,rows,widths,note):blocks.append(('table',(cap,rows,widths,note)))
def fig(name,cap):blocks.append(('figure',(name,cap)))

title('AI导购理由撤回与目标修订仿真研究')
p('本研究通过合成机制仿真，考察重要购买理由撤回后，关联说明与重选操作如何影响目标重评、目标执行及处理负担。分析重点是事实更正与目标处理之间的条件关系，并以无效、反转和已有方法等价的情形检验干预范围。')
h('1 仿真研究方法')
h('模型起点与材料身份',2)
p('仿真以中文人工智能（artificial intelligence，AI）购物助手为情境，固定消费者为iPhone 15选购一只充电器、已有可用连接线、初始预算150元，当前记录预算200元。模型从重要安全理由被撤回后的状态开始，先前理由如何造成预算变化作为情景输入。Gonçalves等（2025）对撤回信息的分析提示，先前证据失效与支持相反结论的新证据并不等价。因此，模型分别处理证据状态、目标关联和目标修订，不以理由撤回自动宣告新预算无效。')
p('模型设置三类背景：依赖已撤回安全理由而提高预算、原先已有独立颜色理由、撤回后已经知情重申。第三类保留既有理解状态，不因未读当前提示而重新变成不理解。三类等量用于覆盖边界，不对应消费者人口比例。商品、价格、偏好及行为参数均为合成设定，未用元宝对话或既有论文效应量估计数值。')
p('合成目录包含79、109、149和179元四件适配商品，以及69元的一件不适配压力项。适配商品的日常充电表现与基础安全设为相同，179元款为蓝色，另设便携、折叠属性。合成任务效用按下式计算；不适配品另设惩罚，并单独报告不适配率。')
eq('utility')
p('式中，i索引合成档案，j索引商品，p为价格，b、c、f分别为蓝色、便携和折叠属性的价值。理由依赖类的颜色价值取0至80，独立理由类和知情重申类取110至200，使合理的升级选择能在模型中保留。该效用用于比较给定偏好下的任务结果，不作为完整自主性的评分。')
h('实验条件与行为过程',2)
p('模型采用商业关系披露D、理由与目标的关联说明R、重选操作C三个二值因素，形成八个组合。主分析固定D为1，比较R与C的四种组合；商业披露与操作的原四组比较固定R为0，保留在补充结果中。另加入全员再确认和强选择性再确认，共十种策略。强选择性基线与所提方案采用相同触发、信息效力和成本，用于检查相同机制能否由已有再确认实现。')
p('关联说明根据公开理由历史触发；基准识别敏感度为0.9、误触发率为0。固定处理量由事实说明、商业披露、操作菜单及实际显示的提示组成：')
eq('load')
p('H为固定处理要求，r为实际提示是否显示。基准商业披露负担为0.4，菜单与关联提示成本各为0.5，注意资源为6。模型将理解更正、注意目标关联、愿意付出操作成本重评及系统执行修订分开。有限注意规则下，更正理解概率基准为0.8，关联线索效力为0.5；均为待校准假设，敏感性分析包含效力为零和高成本条件。')
p('不提供控制菜单时，用户仍可用自然语言修订目标。普通文本与控制操作的基准成本分别为3和0.8，执行成功概率分别为0.8和0.95，并检验相同执行率、控制退化及操作昂贵的边界。商业披露影响线索与阅读成本，不直接消除排序中的商业权重。')
p('比较三种行为规则：完全知情参考规则直接取得并执行该合成世界的完整信息选择；有限注意和默认黏着规则在有限可见项中比较，并保留原推荐的默认影响；满意即止规则依次选择达到合成门槛的候选。选择数量效应具有异质性（Scheibehenne等，2010），因此不预设更多选项总是有益或总会过载；关联文字的阅读成本仍是本模型另设假设。')
h('规模指标与计算验证',2)
p('正式运行使用MATLAB R2023b及基础函数，固定随机种子20261028。基准每类生成40,000个合成条件，共120,000个；其他场景取同一输入中每类前10,000个。所有策略共用商品、偏好及随机数。54个参数场景覆盖零商业偏向、理解缺失、线索无效、注意饱和、过载、执行退化、识别错误和旧依据残留，并包含提示效力与成本的完整网格。')
p('主要指标为目标重评、未重评仍选择超初始预算的商品、明确请求修订后的预算越限、不适配购买及处理负担。未重评升级只表示尚未重评的状态，不编码为未经授权购买。任务效用差以购买为条件，未购买另列；两策略的效用配对差仅在共同购买交集计算。修改请求数量与请求后的条件执行失败率同时报告，以区分执行可靠性和请求暴露量。')
p('不确定性以配对蒙特卡洛标准误（Monte Carlo standard error，MCSE）表示，只衡量模拟积分精度。计算检查涵盖指标范围、预算和兼容边界、同输入重算及强基线等价；另用Python标准库对基准理由依赖类全部40,000条原始记录、三种策略的五项指标及五项配对差重新计算。完整参数、原始输入、代码和诊断表作为配套复现材料保存。')

h('2 仿真结果')
h('关联说明与重选操作的基准结果',2)
p('表1报告商业披露存在、有限注意和默认黏着规则下理由依赖类的结果。四种策略均提供相同事实更正，但目标重评比例不同。仅事实更正时，完成重评的比例为13.47%；同时提供关联说明和重选操作时为51.70%。')
tab('表1 基准情景的配对合成结果',table1,[1.9,.82,.9,.9,.92,.86],
    '注 每组为同一批40,000个理由依赖类合成条件。百分比四舍五入保留两位，差值由未舍入数据计算。处理负担为合成单位，各组未购买比例均为0。')
p(f'固定重选操作可用性后，增加关联说明使完成重评的比例增加{fmt(increase)}个百分点，使未重评仍选择升级款的比例从{rate(b,"dependent_unreviewed_upgrade")}%降至{rate(n,"dependent_unreviewed_upgrade")}%，减少{fmt(drop)}个百分点，配对MCSE为{fmt(se,3)}个百分点。约95%的模拟积分误差带为{fmt(drop-1.96*se)}至{fmt(drop+1.96*se)}个百分点。')
fig('图1_基准机制比较.png','图1 关联说明与重选操作的主指标比较 注 商业披露固定为有，基准有限注意规则，理由依赖类每组40,000个配对合成条件。误差条为1.96倍MCSE，表示模拟精度。')
p(f'同一比较中，不适配购买比例由{rate(b,"noncompatible_purchase")}%升至{rate(n,"noncompatible_purchase")}%，增加{fmt(bad)}个百分点，说明重评覆盖与商品适配是不同维度。四组明确请求预算修订后的越限比例均为0；基准结果因此不提供减少预算越限的差异证据。')
p('修改请求数量由7,819增加至20,678，请求后的条件执行失败率分别为5.24%和5.02%。两组控制执行成功概率均设为0.95，不能将全样本失败数量的变化直接解释为可靠性变化。购买条件下的合成任务效用差由39.644降至37.173；在仅保留兼容购买者的诊断中为24.196和19.556。总效用差对不适配惩罚敏感，适配、未购买和效用结果据此分别报告。')
h('成本边界与反转',2)
tab('表2 预定边界条件及其解释',table2,[1.28,2.55,2.47],
    '注 除完全知情参考与强基线条目外，均固定D为1、C为1，在理由依赖类有限注意规则下比较有无关联说明。结果不代表真实人群发生率。强选择性基线的相同结果来自同构、等参及同随机输入。')
p('在有限注意与满意即止两类规则中，共80个规则与网格组合出现39个正向、31个反向及10个零差异结果。图2显示，关联线索效力和提示成本共同决定方向。在注意资源为6、关联效力为0、提示成本为6的预定组合中，未重评升级增加40.95个百分点，MCSE为0.492个百分点。这是无信息增益却占用资源的模型反例，不是对所有高成本情形的估计。')
fig('图2_条件与反转边界.png','图2 关联线索效力和提示成本的边界 注 固定重选操作可用，比较有无关联说明。红色表示未重评升级减少，蓝色表示增加；数值为百分点。各非基准场景每类10,000个配对合成条件。横轴为离散参数网格。')
p('在信息过载边界中，处理负担由4.916增至7.628，而未重评升级仅由97.98%变为97.92%，该项变化相对配对MCSE较小。零商业偏向时，任务效用差从32.111增至34.837，表明更频繁重评并不保证该合成任务结果改善。')
h('合理升级与已有方法基线',2)
p('独立理由类与知情重申类的颜色价值被设定为足以支持179元选择，因此保留升级不被计为失败。正常触发条件下，已知情重申类不再收到关联提示；全员确认仍显示提示。该比较刻画策略触发属性，不测量真实反感或自主感。')
p('全员再确认在基准理由依赖类的未重评升级比例为29.87%，低于所提方案的31.25%，因为全员提示也覆盖了识别器漏掉的条件。强选择性再确认则在全部场景和规则下与所提方案完全相同。这一结构等价结果说明，同一机制能够通过较强的既有再确认实现，不能据此主张新算法优势。')

h('3 讨论与研究边界')
p('仿真刻画了事实更正与目标处理之间的一条条件性桥梁：当理解、历史关联注意和操作执行存在缺口时，更正商品断言未必完成目标重评。新增关联说明可在有信息效力且成本较低时促进重评，但其对适配、效用与负担的影响需要分别评价。由此，将“是否纠正原句”扩展为“该理由是否公开支持过目标修订”具有可分析的研究对象。')
p('这一机制分析与程序性自主的规范论证共同指向有条件的治理要求：重要理由失去支持、与目标修订有公开联系且决定仍可调整时，说明关联并保留自愿重选机会。说明强度应与重要性和负担相称；已有回复已充分提供信息和机会，或消费者已经知情重申时，不再附加重复确认。措施既允许消费者恢复原预算，也允许其基于独立理由保留升级。')
p('模型使用的关联效力、注意资源、操作成本、执行概率和任务效用均未经真实消费者校准。关联说明有作用是被扫描的机制假设，不能以模拟结果反过来证明该假设在现实成立。完全知情零差异与强基线等价来自参考设定；不适配惩罚、购买条件样本及未购买也影响效用解释。模拟条件数量和MCSE不用于推断中国消费者的总体效果。')
p('研究贡献因此定位为纠错责任范围和触发边界的规范推导，以及可复现的条件机制分析。真实用户是否理解关键取舍、是否认可目标变化，以及说明和操作的净负担，仍需匹配任务或用户研究验证。仿真提供适用条件和反例，不能单独证明伦理义务、首创性或普遍有效性。')
h('本节参考文献')
p('Gonçalves, D., Libgober, J., & Willis, J. (2025). Retractions: Updating from complex information (arXiv:2106.11433v4). arXiv. https://arxiv.org/abs/2106.11433v4')
p('Scheibehenne, B., Greifeneder, R., & Todd, P. M. (2010). Can there ever be too many options? A meta-analytic review of choice overload. Journal of Consumer Research, 37(3), 409–425. https://doi.org/10.1086/651235')

# Put interpretation before the first figure so a large inline visual does not leave a page gap.
fi=next(i for i,b in enumerate(blocks) if b[0]=='figure' and b[1][0].startswith('图1'))
fb=blocks.pop(fi);blocks.insert(fi+2,fb)

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72)
sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
def font(style,size,bold=False,east='SimSun'):
    st=doc.styles[style];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),east)
    return st
normal=font('Normal',12);normal.paragraph_format.line_spacing=1.35;normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.first_line_indent=Pt(24)
font('Title',19,True,'Microsoft YaHei');font('Heading 1',14,True,'Microsoft YaHei');font('Heading 2',12,True,'Microsoft YaHei')
font('Caption',10.5);font('Header',9);font('Footer',9)
for name in ['Title','Heading 1','Heading 2']:
    st=doc.styles[name];st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.keep_with_next=True
    st.paragraph_format.space_before=Pt(12 if name!='Title' else 0);st.paragraph_format.space_after=Pt(6)
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
cap=doc.styles['Caption'];cap.paragraph_format.first_line_indent=Pt(0);cap.paragraph_format.line_spacing=1.15
header=sec.header.paragraphs[0];header.text='AI导购理由撤回与目标修订';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title='AI导购理由撤回与目标修订仿真研究'
doc.core_properties.subject='合成机制仿真的方法结果与讨论';doc.core_properties.comments=''
md=[]
def plainpara(t,style=None):
    z=doc.add_paragraph(t,style);z.paragraph_format.widow_control=True
    return z
def mr(text):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
def sub(base,index):
    z=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(base));s=OxmlElement('m:sub');s.append(mr(index));z.append(e);z.append(s);return z
def mathpara(which):
    z=doc.add_paragraph();z.paragraph_format.first_line_indent=Pt(0);z.alignment=WD_ALIGN_PARAGRAPH.CENTER
    om=OxmlElement('m:oMathPara');o=OxmlElement('m:oMath')
    if which=='utility':
        seq=[sub('V','ij'),mr(' = 250 − '),sub('p','j'),mr(' + '),sub('b','i'),sub('Blue','j'),mr(' + '),sub('c','i'),sub('Compact','j'),mr(' + '),sub('f','i'),sub('Fold','j')]
        tex=r'V_{ij}=250-p_j+b_i Blue_j+c_i Compact_j+f_i Fold_j'
    else:
        seq=[sub('H','i'),mr(' = 1 + 0.4D + '),sub('a','C'),mr(' C + '),sub('a','R'),mr(' '),sub('r','i')]
        tex=r'H_i=1+0.4D+a_C C+a_R r_i'
    for item in seq:o.append(item)
    om.append(o);z._p.append(om);return tex
def maketable(cap,rows,widths,note):
    cp=plainpara(cap,'Caption');cp.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(rows[0]));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
        if i==0:
            rep=OxmlElement('w:tblHeader');trpr.append(rep)
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.text=str(text)
            pr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                v=OxmlElement('w:'+side);v.set(qn('w:w'),'90');v.set(qn('w:type'),'dxa');marg.append(v)
            pr.append(marg)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');pr.append(sh)
            for pp in cell.paragraphs:
                pp.paragraph_format.first_line_indent=Pt(0);pp.paragraph_format.line_spacing=1.15;pp.paragraph_format.space_after=Pt(2)
                for run in pp.runs:
                    run.font.size=Pt(10);run.font.bold=i==0;run.font.color.rgb=RGBColor(0,0,0)
    borders=OxmlElement('w:tblBorders')
    for key in ['top','bottom','insideH']:
        line=OxmlElement('w:'+key);line.set(qn('w:val'),'single');line.set(qn('w:sz'),'6');line.set(qn('w:color'),'888888');borders.append(line)
    t._tbl.tblPr.append(borders)
    nn=plainpara(note,'Caption');nn.paragraph_format.space_after=Pt(8)

for kind,obj in blocks:
    if kind=='title':plainpara(obj,'Title');md+=['# '+obj,'']
    elif kind.startswith('h'):
        lev=int(kind[1:]);plainpara(obj,'Heading '+str(lev));md+=['#'*(lev+1)+' '+obj,'']
    elif kind=='p':
        z=plainpara(obj)
        if obj.startswith(('Gonçalves, D.','Scheibehenne, B.')):
            z.paragraph_format.left_indent=Pt(18);z.paragraph_format.first_line_indent=Pt(-18)
            z.paragraph_format.line_spacing=1.15;z.paragraph_format.keep_together=True
        md+=[obj,'']
    elif kind=='eq':tex=mathpara(obj);md+=['$$',tex,'$$','']
    elif kind=='table':
        cap,rows,widths,note=obj;maketable(cap,rows,widths,note)
        md+=[cap,'','|'+'|'.join(rows[0])+'|','|'+'|'.join(['---']*len(rows[0]))+'|']
        md+=['|'+'|'.join(row)+'|' for row in rows[1:]]+[note,'']
    else:
        name,cap=obj;pp=doc.add_paragraph();pp.paragraph_format.first_line_indent=Pt(0);pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        pp.paragraph_format.keep_with_next=True;pp.add_run().add_picture(str(root/'figures'/name),width=Inches(6.65))
        plainpara(cap,'Caption');md+=['!['+cap+'](figures/'+name+')',cap,'']
docx=root/'论文仿真章节_可编辑稿.docx';doc.save(docx)
(root/'draft.md').write_text('\n'.join(md),encoding='utf-8')

mapping={'date':'2026-10-28','source':'corrected simulation results', 'scope':'simulation methods/results/discussion only',
 'evidence':'synthetic, not human outcomes','rounding':'Decimal half-up; differences from unrounded data',
 'sourceFiles':{},'rows':[]}
for name in ['all_summary.csv','paired_link_contrasts.csv','base_diagnostics.csv','all_grid_effects.csv']:
    f=sim/'results'/name;mapping['sourceFiles'][name]=hashlib.sha256(f.read_bytes()).hexdigest()
for p in core:
    mapping['rows'].append({'scenario':'base','rule':'bounded_default','policy':p,'type':1,
      **{k:get(p)[k] for k in ['n','goal_reviewed','dependent_unreviewed_upgrade','budget_violation_after_request','noncompatible_purchase','processing_load']}})
mapping['pairedContrast']=d
mapping['deliverableSHA256']={f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in [docx,root/'draft.md']}
(root/'数值与出处核对.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
how='''# 并入论文的位置与衔接

Word稿为可编辑的仿真方法、结果和讨论，Markdown为同内容文本源；两张图为从已修正CSV重新导出的中文学术图。表1百分比以十进制四舍五入保留两位，差值使用未舍入结果。

建议放入整篇报告的研究方法、案例与仿真分析、讨论与局限三个位置；合稿时统一章节、图表及参考文献编号。仿真不替代前文的程序性自主论证，元宝预试另作材料身份说明。

前文衔接句可用：前述规范分析提出了重要理由撤回后的目标处理要求。为考察其经验桥梁，本研究进一步构建合成机制模型，比较关联说明、重选操作和既有再确认，并检验无效与高负担条件。

结论衔接句可用：仿真表明，所讨论机制的价值取决于理解、注意与操作成本。治理要求宜落实为与重要性和负担相称的说明与自愿重选，消费者已知情重申后停止重复干预。该分析界定了适用条件，而非由模拟数量证明真人自主性收益。

正文数据来源均为simulation_goal_repair_20261028/results内的最新修正版CSV，数值与出处核对.json记录源文件哈希及主表行。全部54场景、10策略、3规则、原始配对输入及代码仍保留在此前仿真复现包。论文正文保留零结果、反转、兼容性代价及强基线等价，不把未重评升级改称未经授权购买。

本节使用两条已实际读取的文献版本；新增正文并未声称相关论文全文均精读。正式论文其他引用与研究背景仍沿用既有材料账，并在合稿时统一格式。
'''
(root/'并入论文说明.md').write_text(how,encoding='utf-8')
print(json.dumps({'docx':str(docx),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes),
 'nativeEquations':len(doc.element.xpath('//m:oMath')),'markdownCharacters':len('\n'.join(md))},ensure_ascii=False))
