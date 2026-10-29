"""Verify the edited Word package and its already-computed domestic evidence."""
from pathlib import Path
import argparse,csv,hashlib,json,re,zipfile
from lxml import etree as ET
from docx import Document
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parent
DATA=ROOT.parent/'domestic_data_reanalysis_20261029'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--qa-pdf',type=Path,required=True);args=ap.parse_args()
    file=ROOT/'AI导购理由撤回后的预算重审_补强稿国内数据版.docx'
    trace=json.loads((ROOT/'修订核对.json').read_text(encoding='utf-8'))
    original=json.loads((ROOT/'input_package_hashes.json').read_text(encoding='utf-8'))
    doc=Document(file);pdf=PdfReader(args.qa_pdf)
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with zipfile.ZipFile(file) as z:
        package={n:z.read(n) for n in z.namelist()}
    x=ET.fromstring(package['word/document.xml']);body=x.find('w:body',ns)
    text=lambda e:''.join(e.xpath('.//w:t/text()',namespaces=ns))
    blocks=list(body);bi=next(i for i,e in enumerate(blocks) if text(e)=='参考文献')
    ai=next(i for i,e in enumerate(blocks) if text(e)=='附录')
    before='\n'.join(text(e) for e in blocks[:bi]);reftexts=[text(e) for e in blocks[bi+1:ai]]
    full='\n'.join(text(e) for e in blocks)
    refs=[int(re.match(r'^\[(\d+)\]',s).group(1)) for s in reftexts if re.match(r'^\[(\d+)\]',s)]
    first=[]
    for n in re.findall(r'\[(\d+)\]',before):
        n=int(n)
        if n not in first:first.append(n)
    allowed={'word/document.xml','word/media/image2.png','word/media/image3.png','docProps/core.xml'}
    checks={
        'same_package_parts':set(package)==set(original),
        'untargeted_parts_byte_identical':all(hashlib.sha256(b).hexdigest()==original[n]['sha256'] for n,b in package.items() if n not in allowed),
        'original_normative_figure_preserved':hashlib.sha256(package['word/media/image1.png']).hexdigest()==original['word/media/image1.png']['sha256'],
        'all_23_references_resolve':refs==list(range(1,24)) and set(map(int,re.findall(r'\[(\d+)\]',full)))==set(refs),
        'citations_in_first_appearance_order':first==list(range(1,24)),
        'five_tables_three_figures':len(doc.tables)==5 and len(doc.inline_shapes)==3,
        'docx_identity':sha(file)==trace['output_sha256'],
        'docx_author_empty':not doc.core_properties.author and not doc.core_properties.last_modified_by,
        'qa_pdf_author_empty':not (pdf.metadata or {}).get('/Author'),
        'qa_twenty_nonempty_pages':len(pdf.pages)==20 and all(len(p.extract_text().strip())>100 for p in pdf.pages),
        'qa_pdf_within_10M':args.qa_pdf.stat().st_size<10_000_000,
        'old_simulation_results_absent':not any(s in full for s in ['合成机制仿真','40,000','120,000','51.70','31.25','MCSE','54场景','all_summary.csv','run_goal_repair_simulation','388项']),
        'no_template_marker':'{{' not in full,
        'normative_scope_preserved':all(s in full for s in ['普通更正已经使关联可理解','独立理由充分','知情重申及退出','决定地位','条件性比较与预算更新分别判断']),
        'new_behavior_not_generated':trace['new_behavior_generated'] is False and trace['main_analysis_rerun'] is False,
        'primary_results_present':all(s in full.replace('−','-') for s in ['-0.160','[-0.542,0.220]','0.186','[0.032,0.341]','-0.199','[-0.646,0.250]','p=0.584']),
        'innovation_evidence_scope_explicit':all(s in full for s in ['本仿真不检验完整伦理方案，而是为方案提供经验背景','核心创新仍由规范推理、配对规范案例和边界条件论证完成','未直接比较按本判准设计的措施与普通事实纠错、一般再确认或元偏好反思的效果']),
        'normative_cases_distinguished_from_human_validation':'配对规范案例用于判断一致性与反例边界，不等同真人实验验证' in full,
    }
    for name,expected in trace['source_results_sha256'].items():checks['unchanged_'+name]=sha(DATA/'results'/name)==expected
    for i,n in [(2,'图2_国内六臂信任评价.png'),(3,'图3_面板差值与交互边界.png')]:
        checks['domestic_figure_'+str(i)]=hashlib.sha256(package['word/media/image'+str(i)+'.png']).hexdigest()==sha(DATA/'results'/n)
    with (DATA/'results/cell_estimates.csv').open(encoding='utf-8-sig',newline='') as f:
        cells=[r for r in csv.DictReader(f) if r['sample']=='all_516' and r['metric']=='trust_mean']
    actual=[[c.text for c in row.cells] for row in doc.tables[2].rows[1:]]
    checks['six_cell_table_matches_csv']=len(actual)==len(cells)==6 and all(
        row[1:]==[a['participants'],f"{float(a['estimate']):.3f}",f"{float(a['sd']):.3f}",f"[{float(a['ci_low']):.3f}, {float(a['ci_high']):.3f}]"] for row,a in zip(actual,cells))
    exposed='\n'.join(b.decode('utf-8',errors='ignore') for n,b in package.items() if n.endswith(('.xml','.rels')))
    checks['no_private_path_account_or_token']=not re.search(r'[A-Za-z]:[\\/]{1,2}Users[\\/]|wxid_[A-Za-z0-9]+|https?://yuanbao\.tencent\.com/chat|gh[pousr]_[A-Za-z0-9]{20,}|zyj123',exposed,re.I)
    checks['no_tracked_changes_or_comments']=not re.search(r'<w:(?:ins|del|commentRangeStart)\b',package['word/document.xml'].decode('utf-8'))
    checks['retired_files_still_absent']=not any((ROOT.parents[1]/i['path']).is_file() for i in json.loads((ROOT.parents[1]/'provenance/removed_simulations_20261029.json').read_text(encoding='utf-8'))['removed_files'])
    report={'date':'2026-10-30','pass':all(checks.values()),'checks':checks,'file':file.name,'bytes':file.stat().st_size,'sha256':sha(file),
            'pages':len(pdf.pages),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'references':len(refs),'title_characters':len(doc.paragraphs[0].text),'abstract_characters':len(doc.paragraphs[3].text),
            'qa_pdf_bytes':args.qa_pdf.stat().st_size,'qa_pdf_sha256':sha(args.qa_pdf),
            'visual_qa':{'all_final_twenty_pages_checked':True,'review':'2026-10-30: final changed pages 7,8,9,10,11,16 individually inspected; the other 14 final PNGs are byte-identical to the previously inspected report render.','result':'No clipping, overlap, missing glyph or empty page observed; original layout retained.'},
            'evidence_boundary':'Domestic archived human trust data; no direct test of budget-repair efficacy; no newly generated consumer behavior.',
            'not_completed':['real CNKI detection','defense PPT verification','final anonymous supplementary package','official submission']}
    (ROOT/'交付核对.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))
    if not report['pass']:raise SystemExit(1)

if __name__=='__main__':main()
