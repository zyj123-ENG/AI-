"""Check the standalone calculation report against already executed outputs."""
from pathlib import Path
import csv
import hashlib
import json
import re
import zipfile
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
DATA = ROOT.parent / 'domestic_data_reanalysis_20261029'
TITLE = '国内二手数据重抽样分析报告'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    text = (ROOT / 'report.md').read_text(encoding='utf-8')
    docx, pdf = (ROOT / (TITLE + ext) for ext in ['.docx', '.pdf'])
    doc, reader = Document(docx), PdfReader(pdf)
    trace = json.loads((ROOT / '生成记录.json').read_text(encoding='utf-8'))
    audit = json.loads((DATA / 'validation/独立复算.json').read_text(encoding='utf-8'))
    meta = json.loads((DATA / 'results/run_metadata.json').read_text(encoding='utf-8'))
    numeric_text = text.replace('−', '-')
    checks = {
        'only_calculation_report': all(x in text for x in ['变量与统计问题', '回归与重抽样计算', '已运行结果'])
        and not any(x in text for x in ['消费者权益保护法', '## 治理', '## 理论框架']),
        'four_tables_two_figures': len(doc.tables) == 4 and len(doc.inline_shapes) == 2,
        'domestic_observed_sample': meta['participants'] == 516 and meta['campus_sensitivity_n'] == 403
        and meta['new_behavior_generated'] is False and meta['bootstrap_repetitions'] == 5000,
        'independent_point_verification': audit['within_scope_pass'] is True,
        'primary_results_present': all(x in numeric_text for x in
            ['-0.160', '[-0.542,0.220]', '0.186', '[0.032,0.341]', '-0.199', '[-0.646,0.250]', 'p=0.584']),
        'no_placeholders': '{{' not in text and '作者原CSV路径' not in text,
        'six_pages_no_empty_page': len(reader.pages) == 6 and all(len(p.extract_text().strip()) > 100 for p in reader.pages),
        'pdf_within_10M': pdf.stat().st_size < 10_000_000,
        'docx_author_empty': not doc.core_properties.author and not doc.core_properties.last_modified_by,
        'pdf_author_empty': not (reader.metadata or {}).get('/Author'),
        'fixed_original_csv': sha(DATA / 'source/author_original/data/study4_memory_data.csv')
        == 'dec91f502f9fb40f2f89e9e980e981edcf1670e928a91d911a60464c815b0360',
    }
    for name, expected in trace['input_results_sha256'].items():
        checks['unchanged_input_' + name] = sha(DATA / 'results' / name) == expected
    for src, dst in [('图2_国内六臂信任评价.png', '图1_国内六臂信任评价.png'),
                     ('图3_面板差值与交互边界.png', '图2_面板差值与交互边界.png')]:
        checks['original_matlab_figure_' + dst] = sha(DATA / 'results' / src) == sha(ROOT / 'figures' / dst)
    with zipfile.ZipFile(docx) as z:
        xml = '\n'.join(z.read(n).decode('utf-8') for n in z.namelist() if n.endswith(('.xml', '.rels')))
        checks['no_tracked_edits_or_comments'] = not re.search(r'<w:(?:ins|del|commentRangeStart)\b', xml)
    exposed = xml + '\n'.join(p.extract_text() for p in reader.pages)
    checks['no_private_identifiers'] = not re.search(
        r'[A-Za-z]:[\\/]{1,2}Users[\\/]|wxid_[A-Za-z0-9]+|https?://yuanbao\.tencent\.com/chat|gh[pousr]_[A-Za-z0-9]{20,}|zyj123', exposed, re.I)
    # All six displayed means and SDs are rounded from the exact source summary.
    with (DATA / 'results/cell_estimates.csv').open(encoding='utf-8-sig', newline='') as f:
        cells = [x for x in csv.DictReader(f) if x['sample'] == 'all_516' and x['metric'] == 'trust_mean']
    checks['six_cell_means_and_sd_match'] = len(cells) == 6 and all(
        f"{float(x['estimate']):.3f}" in text and f"{float(x['sd']):.3f}" in text for x in cells)
    report = {
        'date': '2026-10-29', 'pass': all(checks.values()), 'checks': checks,
        'scope': 'Standalone simulation/reanalysis section only; no full competition report written by this packaging step.',
        'pages': len(reader.pages), 'cjk': len(re.findall('[\u4e00-\u9fff]', text)),
        'files': {p.name: {'bytes': p.stat().st_size, 'sha256': sha(p)} for p in [docx, pdf]},
        'visual_qa': {'all_final_six_pages_reviewed': True,
                      'pages_1_and_2': 'Final PNGs byte-identical to individually inspected preceding render.',
                      'pages_3_to_6': 'Final PNGs individually reopened after formatting and reproduction-command repairs.',
                      'result': 'No clipped text, overlap, empty page or missing glyph; references fit with content on page 6.'},
        'computation_scope': 'Existing domestic MATLAB OLS and bootstrap outputs reused; independent point/count/F/p validation retained, no independent bootstrap rerun.',
    }
    (ROOT / '交付核对.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if not report['pass']:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
